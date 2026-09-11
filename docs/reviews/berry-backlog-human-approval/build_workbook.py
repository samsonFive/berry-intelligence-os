#!/usr/bin/env python3
"""Build the empty (scope-blocked) berry backlog human-approval workbook.

This script deliberately does not invent backlog rows. It records what was
inspectable in this checkout and leaves Review empty until a live export is
supplied.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = Path(__file__).resolve().parent
TODAY = date.today().isoformat()

REVIEW_COLUMNS = [
    ("record_id", 46),
    ("queue_type", 28),
    ("current_status", 18),
    ("title", 55),
    ("publisher", 24),
    ("source_url", 48),
    ("published_date", 14),
    ("captured_date", 14),
    ("date_basis", 18),
    ("berries", 22),
    ("companies", 28),
    ("geographies", 22),
    ("content_state", 22),
    ("short_summary", 60),
    ("relevance_reason", 40),
    ("quality_flags", 36),
    ("duplicate_of", 46),
    ("proposed_batch", 28),
    ("recommended_action", 22),
    ("recommendation_reason", 55),
    ("final_decision", 18),
    ("reviewer_notes", 36),
]

QUEUE_TYPES = [
    "legacy_unvalidated_evidence",
    "pending_publication",
    "source_fidelity_recovery",
    "atomic_evidence_claim",
]

# Queue-appropriate actions for human final_decision. Legacy purge/purge+block
# are documented but NOT offered as executable workbook decisions.
FINAL_DECISIONS = [
    "approve_as_source",
    "reject_as_source",
    "hold",
    "needs_content_recovery",
    "affirm_source_fidelity",
    "reject_source_fidelity",
    "needs_investigation",
    "approve_atomic_claim",
    "reject_atomic_claim",
    "defer",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}1"


def build() -> dict:
    dataset = {
        "dataset_id": "github-master-checked-in-only",
        "inspected_at": TODAY,
        "git_commit": "ec801d538003c93b53679c8bcdcca5b3a3dba505",
        "git_branch_inspected": "master",
        "repo": "github.com/samsonFive/berry-intelligence-os",
        "requested_but_unavailable": {
            "local_path": "C:/Users/Johnny/Downloads/sscanar/berry-intelligence-os-astra-repair",
            "branch": "fix/astra-news-reader",
            "status_file": "artifacts/astra-repair/REVIEW-STATUS.md",
        },
        "queues": {
            "legacy_unvalidated_evidence": {
                "present": True,
                "live_count": 0,
                "location": "data/evidence/*.json where auto_captured=true and validated=false",
                "note": (
                    "Queue mechanism exists via scripts/export_for_review.py + "
                    "apply_review_decisions.py. Checked-in master has 0 matching rows. "
                    "Historical snapshot review/review-backlog-2026-08-06.xlsx had 1585 "
                    "rows already fully decided (1139 validate / 446 purge) and applied; "
                    "it is NOT treated as the current backlog."
                ),
            },
            "pending_publication": {
                "present": False,
                "live_count": None,
                "location": "inbox/evidence/*.json (gitignored runtime drafts)",
                "note": (
                    "No inbox/ directory in this checkout. Feature-branch code supports "
                    "Promote/Save/Dismiss/Reject on pending drafts. Production backlog "
                    "is not in git."
                ),
            },
            "source_fidelity_recovery": {
                "present": False,
                "live_count": None,
                "location": "inbox/source_fidelity/artifacts/*.json",
                "note": (
                    "No fidelity artifacts present. Supported decisions on feature "
                    "branches: affirmed / rejected / needs_investigation. Affirmation "
                    "does not create facts or authorize extraction by itself beyond "
                    "attaching recovered body text for later extraction readiness."
                ),
            },
            "atomic_evidence_claim": {
                "present": False,
                "live_count": None,
                "location": "inbox drafts with evidence_role=atomic_evidence (+ trusted published atomic)",
                "note": (
                    "No atomic draft queue present. Approving a publication/source must "
                    "NOT auto-approve atomic claims."
                ),
            },
        },
        "records": [],
        "batches": [],
        "exceptions": [
            {
                "exception_id": "EX-DATASET-UNAVAILABLE",
                "severity": "missing_live_backlog",
                "detail": (
                    "Live pending publication, source-fidelity, and atomic queues were "
                    "not available in this Cloud Agent checkout. No article recommendations "
                    "were produced."
                ),
                "required_export": "See HANDOFF-REPORT.md §Exports required",
            },
            {
                "exception_id": "EX-ASTRA-CONTEXT-MISSING",
                "severity": "missing_repair_context",
                "detail": (
                    "Branch fix/astra-news-reader and artifacts/astra-repair/REVIEW-STATUS.md "
                    "are not on origin; not inspected."
                ),
                "required_export": "Push or upload the astra-repair workspace / REVIEW-STATUS.md",
            },
            {
                "exception_id": "EX-DO-NOT-USE-2026-08-06-SNAPSHOT",
                "severity": "stale_decided_snapshot",
                "detail": (
                    "review/review-backlog-2026-08-06.xlsx is a completed historical "
                    "decision file, not an open backlog. Do not re-approve or re-apply it."
                ),
                "required_export": None,
            },
        ],
        "limitations": [
            "No production runtime/inbox data mounted in this environment.",
            "No public source URLs were opened because no pending rows were available to resolve.",
            "Recommendations intentionally left empty to avoid inventing approvals.",
            "Application code and records were not modified.",
        ],
    }

    wb = Workbook()

    # --- Instructions ---
    ws_i = wb.active
    ws_i.title = "Instructions"
    instructions = [
        ("Berry Intelligence OS — Human Approval Workbook", ""),
        ("Generated", TODAY),
        ("Dataset actually reviewed", dataset["dataset_id"]),
        ("Git commit inspected", dataset["git_commit"]),
        ("", ""),
        ("PURPOSE", ""),
        (
            "Use this workbook to approve clear groups of articles, inspect exceptions, "
            "and return final_decision values to the primary coding agent for safe, "
            "queue-scoped application. Leave final_decision blank until you decide.",
            "",
        ),
        ("", ""),
        ("HARD BOUNDARIES", ""),
        ("1. Queues stay distinct. Approving an article as a source does NOT approve claims, create facts, or authorize extraction.", ""),
        ("2. Do not populate executable legacy purge / purge+block in final_decision for bulk apply without an explicit separate ops request.", ""),
        ("3. Preserve existing human decisions; never overwrite them.", ""),
        ("4. Preserve record IDs and source URLs exactly.", ""),
        ("5. Never substitute capture date for publication date.", ""),
        ("6. A legitimate article with missing body text is content-recovery, not automatic rejection.", ""),
        ("7. Do not blanket-approve solely because the publisher is reputable.", ""),
        ("", ""),
        ("QUEUE TYPES", ""),
        ("legacy_unvalidated_evidence", "Checked-in auto_captured evidence awaiting validate/purge-style review."),
        ("pending_publication", "Inbox drafts awaiting Promote / Reject / Dismiss as published Evidence sources."),
        ("source_fidelity_recovery", "Recovered body/transcript artifacts awaiting fidelity affirmation for trusted Evidence."),
        ("atomic_evidence_claim", "Atomic claim drafts awaiting approve/reject as factual claims — separate from source approval."),
        ("", ""),
        ("RECOMMENDED_ACTION / FINAL_DECISION MEANINGS", ""),
        ("approve_as_source", "Queue: pending_publication or legacy. Maps to Promote/publish or validate. Marks the article as accepted Evidence/source only."),
        ("reject_as_source", "Queue: pending_publication or legacy. Maps to Reject (pending) or non-executable reject recommendation for legacy. Does not block domains."),
        ("hold", "Leave untouched; needs human judgment beyond the recommendation."),
        ("needs_content_recovery", "Legitimate source candidate with missing/inaccessible body. Route to source-fidelity/content-recovery — do NOT treat as claim approval."),
        ("affirm_source_fidelity", "Queue: source_fidelity_recovery only. Maps to decision=affirmed. Attaches recovered content for extraction readiness; does NOT create facts/claims."),
        ("reject_source_fidelity", "Queue: source_fidelity_recovery only. Maps to decision=rejected."),
        ("needs_investigation", "Queue: source_fidelity_recovery only. Maps to decision=needs_investigation."),
        ("approve_atomic_claim", "Queue: atomic_evidence_claim only. Maps to publish/approve of atomic draft. Never implied by source approval."),
        ("reject_atomic_claim", "Queue: atomic_evidence_claim only. Maps to reject of atomic draft."),
        ("defer", "Park for a later review batch; no state change."),
        ("", ""),
        ("LEGACY SCRIPT MAPPING (informational — not auto-executed from this workbook)", ""),
        ("validate", "scripts/apply_review_decisions.py — sets validated=true, increments source validated_count."),
        ("purge", "DESTRUCTIVE — deletes evidence file. Not populated as an executable recommendation here."),
        ("purge+block", "DESTRUCTIVE — deletes evidence and adds domain to blocklist. Not populated here."),
        ("", ""),
        ("HOW THE PRIMARY AGENT SHOULD APPLY", ""),
        ("1. Read only rows where final_decision is non-blank.", ""),
        ("2. Apply each decision only within that row's queue_type using the mapping above.", ""),
        ("3. Skip any row whose record_id no longer exists; report it.", ""),
        ("4. Never cascade: source approval must not create facts, relationships, or atomic approvals.", ""),
        ("5. Do not run importers or modify production trust tallies beyond the queue-specific transition.", ""),
        ("", ""),
        ("CURRENT FILL STATUS", ""),
        ("Review tab rows", "0 — live backlog unavailable in this environment"),
        ("Batches proposed", "0"),
        ("Exceptions logged", str(len(dataset["exceptions"]))),
        ("Next step", "Supply the exports listed on the Exceptions tab / HANDOFF-REPORT.md, then re-run review."),
    ]
    ws_i["A1"] = "Item"
    ws_i["B1"] = "Detail"
    for cell in ws_i[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for i, (a, b) in enumerate(instructions, start=2):
        ws_i.cell(i, 1, a)
        ws_i.cell(i, 2, b)
        if a in {"PURPOSE", "HARD BOUNDARIES", "QUEUE TYPES", "RECOMMENDED_ACTION / FINAL_DECISION MEANINGS",
                 "LEGACY SCRIPT MAPPING (informational — not auto-executed from this workbook)",
                 "HOW THE PRIMARY AGENT SHOULD APPLY", "CURRENT FILL STATUS"}:
            ws_i.cell(i, 1).font = Font(bold=True, color="1F4E79")
    ws_i.column_dimensions["A"].width = 42
    ws_i.column_dimensions["B"].width = 110
    ws_i.freeze_panes = "A2"

    # --- Review ---
    ws_r = wb.create_sheet("Review")
    for idx, (name, width) in enumerate(REVIEW_COLUMNS, start=1):
        cell = ws_r.cell(1, idx, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws_r.column_dimensions[get_column_letter(idx)].width = width
    ws_r.freeze_panes = "A2"
    ws_r.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_COLUMNS))}1"
    # Data validations ready for when rows are added
    queue_dv = DataValidation(type="list", formula1='"' + ",".join(QUEUE_TYPES) + '"', allow_blank=True)
    action_dv = DataValidation(type="list", formula1='"' + ",".join(FINAL_DECISIONS) + '"', allow_blank=True)
    ws_r.add_data_validation(queue_dv)
    ws_r.add_data_validation(action_dv)
    queue_dv.add("B2:B5000")
    action_dv.add("S2:S5000")  # recommended_action
    action_dv.add("U2:U5000")  # final_decision
    # Placeholder note row so filters/freeze are obvious; clearly not a record
    note = {
        "record_id": "(no live backlog rows in this dataset)",
        "queue_type": "",
        "current_status": "unavailable",
        "title": "See Exceptions + HANDOFF-REPORT.md — do not treat this placeholder as an evidence ID",
        "publisher": "",
        "source_url": "",
        "published_date": "",
        "captured_date": "",
        "date_basis": "",
        "berries": "",
        "companies": "",
        "geographies": "",
        "content_state": "not_inspected",
        "short_summary": "Live queues were not present in the Cloud Agent checkout.",
        "relevance_reason": "",
        "quality_flags": "DATASET_UNAVAILABLE",
        "duplicate_of": "",
        "proposed_batch": "",
        "recommended_action": "",
        "recommendation_reason": "No recommendation issued because no pending records were available.",
        "final_decision": "",
        "reviewer_notes": "",
    }
    ws_r.append([note[name] for name, _ in REVIEW_COLUMNS])
    for col in range(1, len(REVIEW_COLUMNS) + 1):
        ws_r.cell(2, col).fill = WARN_FILL

    # --- Batches ---
    ws_b = wb.create_sheet("Batches")
    batch_headers = [
        ("proposed_batch", 28),
        ("queue_type", 28),
        ("row_count", 12),
        ("review_rationale", 80),
        ("suggested_action", 24),
        ("includes_exceptions", 18),
        ("status", 18),
    ]
    for idx, (name, width) in enumerate(batch_headers, start=1):
        cell = ws_b.cell(1, idx, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws_b.column_dimensions[get_column_letter(idx)].width = width
    ws_b.freeze_panes = "A2"
    ws_b.append([
        "(none)",
        "n/a",
        0,
        "No batches proposed: live backlog unavailable. After export is supplied, batches should group by shared rationale (e.g. clear commercial berry trade press; false-match Raspberry Pi; cookie-wall recovery cases) with exceptions separated.",
        "",
        "n/a",
        "blocked_on_export",
    ])

    # --- Exceptions ---
    ws_e = wb.create_sheet("Exceptions")
    ex_headers = [
        ("exception_id", 28),
        ("severity", 28),
        ("detail", 90),
        ("required_export", 70),
    ]
    for idx, (name, width) in enumerate(ex_headers, start=1):
        cell = ws_e.cell(1, idx, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws_e.column_dimensions[get_column_letter(idx)].width = width
    ws_e.freeze_panes = "A2"
    for ex in dataset["exceptions"]:
        ws_e.append([ex["exception_id"], ex["severity"], ex["detail"], ex.get("required_export") or ""])
        for col in range(1, 5):
            ws_e.cell(ws_e.max_row, col).fill = WARN_FILL

    xlsx_path = OUT_DIR / "berry-backlog-review.xlsx"
    json_path = OUT_DIR / "berry-backlog-review.json"
    wb.save(xlsx_path)
    json_path.write_text(json.dumps(dataset, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return {"xlsx": str(xlsx_path), "json": str(json_path), "dataset": dataset}


if __name__ == "__main__":
    result = build()
    print(json.dumps({k: result[k] for k in ("xlsx", "json")}, indent=2))
